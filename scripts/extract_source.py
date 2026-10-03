"""Read-only, deterministic extraction. No workbook formulas are executed as code."""

import ast, operator, re, json, hashlib, shutil
from decimal import Decimal
from pathlib import Path
from collections import defaultdict
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / "backend" / "data"
DEST.mkdir(parents=True, exist_ok=True)
path = Path.home() / "Downloads" / "Auto Qoutation File.xlsx"
wb = openpyxl.load_workbook(path)


def numeric(ws, ref, seen=None):
    seen = set() if seen is None else seen.copy()
    if ref in seen:
        raise ValueError("Circular reference " + ref)
    seen.add(ref)
    v = ws[ref].value
    if v is None:
        return Decimal(0)
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    if not isinstance(v, str) or not v.startswith("="):
        raise ValueError("Not numeric " + str(v))
    expr = re.sub(
        r"\$?([A-Z]+)\$?(\d+)", lambda m: str(numeric(ws, m[1] + m[2], seen)), v[1:]
    )

    def visit(n):
        if isinstance(n, ast.Constant) and isinstance(n.value, (int, float)):
            return Decimal(str(n.value))
        if isinstance(n, ast.UnaryOp) and isinstance(n.op, (ast.UAdd, ast.USub)):
            return visit(n.operand) * (1 if isinstance(n.op, ast.UAdd) else -1)
        if isinstance(n, ast.BinOp) and type(n.op) in (
            ast.Add,
            ast.Sub,
            ast.Mult,
            ast.Div,
        ):
            return {
                ast.Add: operator.add,
                ast.Sub: operator.sub,
                ast.Mult: operator.mul,
                ast.Div: operator.truediv,
            }[type(n.op)](visit(n.left), visit(n.right))
        raise ValueError("Unsupported formula " + v)

    return visit(ast.parse(expr, mode="eval").body)


ws = wb["Master List"]
products = []
item = ""
groups = defaultdict(list)
issues = []
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 1).value:
        item = str(ws.cell(r, 1).value).strip()
    if not any(ws.cell(r, c).value is not None for c in range(2, 9)):
        continue
    raw = {ws.cell(1, c).value: ws.cell(r, c).value for c in range(1, 9)}
    if not item:
        issues.append({"row": r, "reason": "Missing item", "raw": raw})
        continue
    key = [item] + [str(ws.cell(r, c).value or "-").strip() for c in range(2, 5)]
    price = None
    try:
        if ws.cell(r, 6).value is not None:
            price = str(numeric(ws, "F" + str(r)))
    except (ValueError, ArithmeticError) as e:
        issues.append({"row": r, "reason": str(e)})
    product = dict(zip(["item", "carcass", "shutter", "finish"], key))
    product.update(
        source_row=r,
        price=price,
        price_display=price or "On Request",
        other=str(numeric(ws, "E" + str(r))),
        hardware=str(numeric(ws, "G" + str(r))),
        pricing_area=str(numeric(ws, "H" + str(r))),
        formula=ws.cell(r, 6).value if ws.cell(r, 6).data_type == "f" else None,
        raw=raw,
    )
    groups["|".join(key).casefold()].append(r)
    products.append(product)
flat = {str(r[5]).casefold(): r for r in list(wb["MasterFlat"].values)[1:]}
for p in products:
    k = "|".join(p[f] for f in ["item", "carcass", "shutter", "finish"]).casefold()
    f = flat.get(k)
    if f is None:
        issues.append({"row": p["source_row"], "reason": "Missing from MasterFlat"})
    elif isinstance(f[4], (int, float)) and Decimal(str(f[4])) != Decimal(
        p["price"] or 0
    ):
        issues.append(
            {
                "row": p["source_row"],
                "reason": "Master List / MasterFlat price conflict",
                "master": p["price"],
                "flat": f[4],
            }
        )
    elif p["price"] is None and f:
        p["price_display"] = str(f[4])
duplicates = [{"combination": k, "rows": rs} for k, rs in groups.items() if len(rs) > 1]
ws = wb["Auto Qoutation"]
sections = []
section = None
for r in range(13, 68):
    h = ws.cell(r, 8).value
    if isinstance(h, str) and h != "TOTAL":
        section = {"name": h, "source_row": r, "items": []}
        sections.append(section)
    if isinstance(h, (int, float)):
        keys = [str(ws.cell(r, c).value or "").strip() for c in range(1, 5)]
        key = "|".join(keys).casefold()
        matches = [
            p
            for p in products
            if "|".join(
                p[f] for f in ["item", "carcass", "shutter", "finish"]
            ).casefold()
            == key
        ]

        def measure(c):
            try:
                return str(numeric(ws, c + str(r)))
            except ValueError:
                return None

        desc = ws.cell(r, 10).value
        if isinstance(desc, str) and desc.startswith("="):

            def celltext(m):
                return json.dumps(str(ws[m[0]].value or ""), ensure_ascii=False)

            parts = re.findall(r'"[^"]*"|[A-Z]+\d+', desc[1:])
            desc = "".join(
                p[1:-1] if p.startswith('"') else str(ws[p].value or "") for p in parts
            )
        row = {
            "source_row": r,
            "source_serial": h,
            "product_source_row": (
                matches[0]["source_row"] if len(matches) == 1 else None
            ),
            "item_label": ws.cell(r, 9).value or (keys[0] or ""),
            "description": desc or "",
            "width": measure("K"),
            "length": measure("L"),
            "quantity": measure("N"),
            "measurement_mode": "rft" if ws.cell(r, 11).value == "Rft" else "area",
            "manual_area": measure("M") if ws.cell(r, 11).value == "Rft" else None,
            "other_description": ws.cell(r, 6).value or "",
            "other_amount": measure("G") if r == 31 else "0",
            "flat_charge": "4000" if r == 56 else "0",
            "raw_selection": keys,
            "source_manual_rate": (
                ws.cell(r, 15).value
                if isinstance(ws.cell(r, 15).value, (int, float))
                else (
                    ws.cell(r, 5).value
                    if isinstance(ws.cell(r, 5).value, (int, float))
                    else None
                )
            ),
        }
        if not row["product_source_row"] or matches and matches[0]["price"] is None:
            issues.append(
                {
                    "quotation_row": r,
                    "reason": "Product selection or numeric price requires review",
                    "selection": keys,
                    "manual_rate": row["source_manual_rate"],
                }
            )
        section["items"].append(row)
terms = [ws.cell(r, 9).value for r in range(73, 86)]
data = {
    "source_file": path.name,
    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    "products": products,
    "company_text": ws["H2"].value,
    "terms": terms,
    "section_names": [s["name"] for s in sections],
    "additional_section_names": [
        wb["new format "].cell(r, 1).value for r in [4, 8, 15, 24, 30, 36]
    ],
    "quotation": {
        "customer_name": ws["H9"].value,
        "address": ws["K9"].value,
        "project_reference": ws["O9"].value,
        "sections": sections,
    },
    "print": {
        "orientation": "portrait",
        "size": "A4",
        "margins_inches": {"left": 0.75, "right": 0.75, "top": 1, "bottom": 1},
        "columns": {k: v.width for k, v in ws.column_dimensions.items() if k >= "H"},
        "header_fill": "#B8CCE4",
        "section_fill": "#E9E791",
        "font": "Calibri",
    },
    "report": {
        "duplicates": duplicates,
        "issues": issues,
        "product_count": len(products),
        "flat_count": len(flat),
    },
}
(DEST / "source.json").write_text(
    json.dumps(data, ensure_ascii=False, indent=2), encoding="utf8"
)
shutil.copy2(path, DEST / path.name)
for i, image in enumerate(ws._images):
    (DEST / "logo.png").write_bytes(image._data())
print(json.dumps(data["report"], ensure_ascii=True, indent=2))
