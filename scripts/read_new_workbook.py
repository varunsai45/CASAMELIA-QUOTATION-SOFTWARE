from pathlib import Path
import openpyxl, json, zipfile
root=Path(__file__).resolve().parents[1]
p=Path.home()/'Downloads'/'Qoute_1-10-2026.xlsx'
out=root/'docs'/'new-source'
out.mkdir(parents=True,exist_ok=True)
w=openpyxl.load_workbook(p,read_only=True,data_only=False)
d=openpyxl.load_workbook(p,read_only=True,data_only=True)
result={}
for s in w:
    cells=[{'cell':c.coordinate,'value':c.value,'type':c.data_type,'cached':d[s.title][c.coordinate].value} for row in s.iter_rows() for c in row if c.value is not None]
    result[s.title]=cells
    (out/(s.title.strip()+'.txt')).write_text('\n'.join(f"{c['cell']}: {c['value']} [cached={c['cached']}]" for c in cells),encoding='utf8')
    print(s.title,len(cells),flush=True)
(out/'cells.json').write_text(json.dumps(result,default=str,ensure_ascii=False,indent=2),encoding='utf8')
with zipfile.ZipFile(p) as z:
    for n in z.namelist():
        if n.startswith('xl/') and n.endswith('.xml'):
            (out/n.replace('/','--')).write_bytes(z.read(n))
