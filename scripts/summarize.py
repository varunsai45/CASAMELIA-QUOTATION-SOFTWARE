import openpyxl, collections
from pathlib import Path

wb = openpyxl.load_workbook(Path.home() / "Downloads" / "Auto Qoutation File.xlsx")
cv = openpyxl.load_workbook(
    Path.home() / "Downloads" / "Auto Qoutation File.xlsx", data_only=True
)
for name in ["Master List", "Auto Qoutation", "new format "]:
    w = wb[name]
    print("\nSHEET", name)
    if name == "Master List":
        for row in w:
            if any(c.value is not None for c in row[4:]):
                print(
                    "ROW",
                    row[0].row,
                    [
                        (c.coordinate, c.value, cv[name][c.coordinate].value)
                        for c in row
                        if c.value is not None
                    ],
                )
    else:
        for row in w:
            r = row[0].row
            if r < 14 or r > 75 or name == "new format " and r < 25:
                entries = [
                    (c.coordinate, c.value, cv[name][c.coordinate].value)
                    for c in row
                    if c.value is not None
                    and not (name == "Auto Qoutation" and c.column == 5)
                ]
                if entries:
                    print(entries)
        print(
            "SECTION HEADERS",
            [
                (c.coordinate, c.value)
                for row in w
                for c in row
                if c.column == (8 if name == "Auto Qoutation" else 1)
                and isinstance(c.value, str)
                and not c.value.startswith("=")
            ],
        )
        print(
            "RATE FORMULAS",
            [
                (c.coordinate, c.value)
                for row in w
                for c in row
                if c.column == (15 if name == "Auto Qoutation" else 13)
                and c.data_type == "f"
            ],
        )
        print(
            "STYLES",
            [
                (w[p].coordinate, w[p].style_id, str(w[p].font), str(w[p].fill))
                for p in (
                    ["H2", "H11", "H13", "H14", "P16", "H69"]
                    if name == "Auto Qoutation"
                    else ["A1", "A11", "A13"]
                )
            ],
        )
print("FLAT", cv["MasterFlat"].max_row, list(cv["MasterFlat"].values)[:3])
print(
    "DUPLICATES",
    [
        (k, n)
        for k, n in collections.Counter(
            r[5] for r in list(cv["MasterFlat"].values)[1:]
        ).items()
        if n > 1
    ],
)
