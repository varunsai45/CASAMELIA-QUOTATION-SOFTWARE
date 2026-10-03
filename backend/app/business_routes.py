"""Configurable areas, company settings, price history and reviewed imports."""

import base64, hashlib, json, zipfile
from io import BytesIO
from decimal import Decimal
from PIL import Image
import openpyxl
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy import select, delete
from .models import (
    Area,
    AreaProduct,
    MasterProduct,
    MasterPriceHistory,
    SourcePrice,
    Setting,
    SettingVersion,
    ImportBatch,
    PriceConflict,
    AuditLog,
    User,
)
from .schemas import (
    AreaInput,
    AreaAssignment,
    CompanyInput,
    NumberingInput,
    ImportInput,
)
from .security import current_user, admin
from .db import get_db

router = APIRouter()


def log(db, user, action, entity, before=None, after=None):
    db.add(
        AuditLog(
            actor_id=user.id, action=action, entity=entity, before=before, after=after
        )
    )


def area_dict(a):
    return {
        k: getattr(a, k) for k in ["id", "name", "active", "position", "source_sheet"]
    }


@router.get("/areas")
def areas(
    include_inactive: bool = False, user=Depends(current_user), db=Depends(get_db)
):
    stmt = select(Area)
    if user.role != "admin" or not include_inactive:
        stmt = stmt.where(Area.active.is_(True))
    return [area_dict(a) for a in db.scalars(stmt.order_by(Area.position, Area.id))]


@router.post("/areas", status_code=201)
def add_area(data: AreaInput, user=Depends(admin), db=Depends(get_db)):
    a = Area(**data.model_dump())
    db.add(a)
    db.flush()
    log(db, user, "area_created", f"area:{a.id}", after=area_dict(a))
    db.commit()
    return area_dict(a)


@router.put("/areas/{aid}")
def edit_area(aid: int, data: AreaInput, user=Depends(admin), db=Depends(get_db)):
    a = db.get(Area, aid)
    if not a:
        raise HTTPException(404, "Area not found.")
    before = area_dict(a)
    for k, v in data.model_dump().items():
        setattr(a, k, v)
    log(db, user, "area_updated", f"area:{aid}", before, area_dict(a))
    db.commit()
    return area_dict(a)


@router.get("/areas/{aid}/products")
def assignments(aid: int, user=Depends(admin), db=Depends(get_db)):
    if not db.get(Area, aid):
        raise HTTPException(404, "Area not found.")
    return {
        "product_ids": list(
            db.scalars(select(AreaProduct.product_id).where(AreaProduct.area_id == aid))
        )
    }


@router.put("/areas/{aid}/products")
def update_assignments(
    aid: int, data: AreaAssignment, user=Depends(admin), db=Depends(get_db)
):
    if not db.get(Area, aid):
        raise HTTPException(404, "Area not found.")
    if len(set(data.product_ids)) != len(data.product_ids) or any(
        not db.get(MasterProduct, pid) for pid in data.product_ids
    ):
        raise HTTPException(422, "Select valid, unique products.")
    before = list(
        db.scalars(select(AreaProduct.product_id).where(AreaProduct.area_id == aid))
    )
    db.execute(delete(AreaProduct).where(AreaProduct.area_id == aid))
    for pid in data.product_ids:
        db.add(AreaProduct(area_id=aid, product_id=pid))
    log(
        db,
        user,
        "area_product_assignments",
        f"area:{aid}",
        {"product_ids": before},
        data.model_dump(),
    )
    db.commit()
    return {"status": "saved"}


@router.get("/master-list/{pid}/history")
def history(pid: int, user=Depends(admin), db=Depends(get_db)):
    if not db.get(MasterProduct, pid):
        raise HTTPException(404, "Product not found.")
    return {
        "prices": [
            {
                "old_price": str(h.old_price) if h.old_price is not None else None,
                "new_price": str(h.new_price) if h.new_price is not None else None,
                "reason": h.reason,
                "date": h.created_at.isoformat(),
                "by": (
                    db.get(User, h.actor_id).username if h.actor_id else "source import"
                ),
            }
            for h in db.scalars(
                select(MasterPriceHistory)
                .where(MasterPriceHistory.product_id == pid)
                .order_by(MasterPriceHistory.id.desc())
            )
        ],
        "sources": [
            {
                "sheet": s.source_sheet,
                "row": s.source_row,
                "file": s.source_file,
                "price": str(s.price) if s.price is not None else None,
                "raw": s.raw,
            }
            for s in db.scalars(
                select(SourcePrice).where(SourcePrice.product_id == pid)
            )
        ],
    }


@router.put("/settings/company")
def company(data: CompanyInput, user=Depends(admin), db=Depends(get_db)):
    value = data.model_dump()
    if data.logo_base64:
        try:
            raw = base64.b64decode(data.logo_base64, validate=True)
            im = Image.open(BytesIO(raw))
            if im.format not in ["PNG", "JPEG"] or im.width > 3000 or im.height > 3000:
                raise ValueError()
            im.verify()
        except Exception:
            raise HTTPException(
                422, "Select a valid PNG or JPEG logo up to 3000 pixels."
            )
    s = db.get(Setting, "company")
    before = s.value
    s.value = value
    db.add(SettingVersion(key="company", value=value, actor_id=user.id))
    log(db, user, "company_updated", "company", before, value)
    db.commit()
    return value


@router.put("/settings/numbering")
def numbering(data: NumberingInput, user=Depends(admin), db=Depends(get_db)):
    s = db.get(Setting, "numbering")
    before = s.value if s else None
    if not s:
        s = Setting(key="numbering", value=data.model_dump())
        db.add(s)
    else:
        s.value = data.model_dump()
    log(db, user, "numbering_updated", "numbering", before, s.value)
    db.commit()
    return s.value


@router.get("/settings/history")
def settings_history(user=Depends(admin), db=Depends(get_db)):
    return [
        {
            "id": s.id,
            "key": s.key,
            "value": s.value,
            "date": s.created_at.isoformat(),
            "actor_id": s.actor_id,
        }
        for s in db.scalars(
            select(SettingVersion).order_by(SettingVersion.id.desc()).limit(100)
        )
    ]


def parse_upload(raw):
    try:
        with zipfile.ZipFile(BytesIO(raw)) as z:
            if (
                sum(f.file_size for f in z.infolist()) > 40_000_000
                or len(z.infolist()) > 2000
            ):
                raise ValueError("Workbook is too large when expanded.")
        w = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=False)
        cache = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=True)
        if "Master List" not in w:
            raise ValueError("A Master List worksheet is required.")
        s = w["Master List"]
        cached = {
            c.coordinate: c.value
            for row in cache["Master List"]
            for c in row
            if c.value is not None
        }
        if s.max_row > 20000 or s.max_column > 300:
            raise ValueError("Master List exceeds import limits.")
        rows = list(s.iter_rows(values_only=True))
        header = [str(x or "").strip().casefold() for x in rows[0]]
        if header[:3] != ["item", "specification", "price"]:
            raise ValueError(
                "Expected Master List columns: Item, Specification, Price. Use the October workbook format."
            )
        products = []
        item = ""
        seen = set()
        duplicates = []
        flat = {}
        if "MasterFlat" in w:
            for row in cache["MasterFlat"].iter_rows(min_row=2, values_only=True):
                if len(row) >= 5 and row[0]:
                    key = (
                        str(row[0]).strip().casefold(),
                        " + ".join(
                            str(x).strip() for x in row[1:4] if x and x != "-"
                        ).casefold(),
                    )
                    flat[key] = (
                        str(row[4]) if isinstance(row[4], (int, float)) else None
                    )
        for n, row in enumerate(rows[1:], 2):
            item = str(row[0] or item).strip()
            spec = str(row[1] or "").strip()
            if not spec:
                continue
            price = cached.get("C" + str(n))
            if price is not None and not isinstance(price, (int, float)):
                price = None
            if price is not None and (
                not Decimal(str(price)).is_finite() or price < 0 or price > 100000000
            ):
                raise ValueError(f"Invalid price at row {n}.")
            key = item.casefold() + "|" + spec.casefold()
            if key in seen:
                duplicates.append({"key": key, "row": n})
            seen.add(key)
            lookup = flat.get((item.casefold(), spec.casefold()))
            products.append(
                {
                    "key": key,
                    "item": item,
                    "specification": spec,
                    "price": str(price) if price is not None else None,
                    "lookup_price": lookup,
                    "row": n,
                    "formula": row[2],
                }
            )
        if not products:
            raise ValueError("No products were found.")
        return {"products": products, "duplicates": duplicates, "sheets": w.sheetnames}
    except Exception as e:
        raise HTTPException(
            422,
            (
                str(e)
                if isinstance(e, ValueError)
                else "The uploaded workbook could not be read."
            ),
        )


@router.post("/imports/preview", status_code=201)
def import_preview(data: ImportInput, user=Depends(admin), db=Depends(get_db)):
    try:
        raw = base64.b64decode(data.content_base64, validate=True)
    except Exception:
        raise HTTPException(422, "Invalid file content.")
    if len(raw) > 5_000_000 or not data.filename.lower().endswith(".xlsx"):
        raise HTTPException(422, "Select an XLSX file under 5 MB.")
    parsed = parse_upload(raw)
    existing = {
        p.combination_key.removeprefix("october|"): p
        for p in db.scalars(
            select(MasterProduct).where(
                MasterProduct.catalogue == "october",
                MasterProduct.source_sheet == "Master List",
            )
        )
    }
    report = {
        "new_items": [],
        "updated_items": [],
        "price_changes": [],
        "conflicts": [],
        "duplicates": parsed["duplicates"],
        "removed_items": [],
        "warnings": [
            "Area-sheet changes are preserved in the uploaded file and listed by sheet. This reviewed import updates Master List rows; manage area assignments and add area-only specifications in the Admin catalogue."
        ],
    }
    for row in parsed["products"]:
        p = existing.get(row["key"])
        row["expected_price"] = str(p.price) if p and p.price is not None else None
        row["existing_id"] = p.id if p else None
        if not p:
            report["new_items"].append(row)
        elif row["price"] is not None and (
            p.price is None or Decimal(row["price"]) != p.price
        ):
            report["price_changes"].append(
                {
                    "id": p.id,
                    "item": p.item,
                    "old": str(p.price) if p.price is not None else None,
                    "new": row["price"],
                }
            )
        else:
            report["updated_items"].append(
                {"id": p.id, "item": p.item, "specification": p.specification}
            )
        if (
            row["lookup_price"] is not None
            and row["price"] is not None
            and Decimal(row["lookup_price"]) != Decimal(row["price"])
        ):
            report["conflicts"].append(row)
    keys = {r["key"] for r in parsed["products"]}
    report["removed_items"] = [
        {"id": p.id, "item": p.item, "specification": p.specification}
        for k, p in existing.items()
        if k not in keys and p.active
    ]
    batch = ImportBatch(
        filename=data.filename,
        sha256=hashlib.sha256(raw).hexdigest(),
        report=report,
        data=parsed,
        source_file=raw,
        actor_id=user.id,
    )
    db.add(batch)
    db.flush()
    log(db, user, "import_preview", f"import:{batch.id}", after=report)
    db.commit()
    return {"id": batch.id, "report": report, "status": batch.status}


@router.post("/imports/{bid}/apply")
def apply_import(bid: int, user=Depends(admin), db=Depends(get_db)):
    batch = db.scalar(
        select(ImportBatch).where(ImportBatch.id == bid).with_for_update()
    )
    if not batch or batch.status != "preview":
        raise HTTPException(
            409, "This import preview is unavailable or already applied."
        )
    if batch.data["duplicates"]:
        raise HTTPException(
            422, "Resolve duplicate combinations in the workbook before importing."
        )
    for row in batch.data["products"]:
        p = db.get(MasterProduct, row["existing_id"]) if row["existing_id"] else None
        current = str(p.price) if p and p.price is not None else None
        if (
            p
            and (current is None) != (row["expected_price"] is None)
            or p
            and current is not None
            and Decimal(current) != Decimal(row["expected_price"])
        ):
            raise HTTPException(
                409, "Master prices changed after this preview. Create a new preview."
            )
        if not p:
            parts = row["specification"].split(" + ")
            p = MasterProduct(
                item=row["item"],
                specification=row["specification"],
                carcass=parts[0],
                shutter=parts[1] if len(parts) > 2 else "-",
                finish=" + ".join(parts[2:] if len(parts) > 2 else parts[1:]) or "-",
                category=row["item"],
                combination_key="october|" + row["key"],
                catalogue="october",
                source_sheet="Master List",
                source_row=row["row"],
                source_data=row,
            )
            db.add(p)
            db.flush()
        old = p.price
        conflict = (
            row["lookup_price"] is not None
            and row["price"] is not None
            and Decimal(row["lookup_price"]) != Decimal(row["price"])
        )
        if conflict:
            c = db.scalar(
                select(PriceConflict).where(
                    PriceConflict.product_id == p.id,
                    PriceConflict.status == "unresolved",
                )
            )
            if c and (
                Decimal(c.visible_price) != Decimal(row["price"])
                or Decimal(c.lookup_price) != Decimal(row["lookup_price"])
            ):
                raise HTTPException(
                    409,
                    "Resolve the existing price conflict before importing changed conflicting values. Both sources remain preserved in the import preview.",
                )
            if not c:
                db.add(
                    PriceConflict(
                        product_id=p.id,
                        item=p.item,
                        carcass=p.carcass,
                        shutter=p.shutter,
                        finish=p.finish,
                        visible_price=row["price"],
                        lookup_price=row["lookup_price"],
                    )
                )
            p.price = None
            p.needs_review = True
        elif not p.needs_review:
            p.price = row["price"]
        p.active = True
        db.add(
            SourcePrice(
                product_id=p.id,
                source_file=batch.filename,
                source_sheet="Master List",
                source_row=row["row"],
                price=row["price"],
                raw=row,
            )
        )
        if conflict:
            db.add(
                SourcePrice(
                    product_id=p.id,
                    source_file=batch.filename,
                    source_sheet="MasterFlat",
                    source_row=row["row"],
                    price=row["lookup_price"],
                    raw=row,
                )
            )
        if old != p.price:
            db.add(
                MasterPriceHistory(
                    product_id=p.id,
                    old_price=old,
                    new_price=p.price,
                    actor_id=user.id,
                    reason=f"Reviewed import {batch.id}",
                )
            )
    for row in batch.report["removed_items"]:
        p = db.get(MasterProduct, row["id"])
        p.active = False
    batch.status = "applied"
    log(db, user, "import_applied", f"import:{bid}", after=batch.report)
    db.commit()
    return {"status": "applied", "report": batch.report}


@router.get("/imports")
def imports(user=Depends(admin), db=Depends(get_db)):
    return [
        {
            "id": b.id,
            "filename": b.filename,
            "status": b.status,
            "report": b.report,
            "date": b.created_at.isoformat(),
        }
        for b in db.scalars(
            select(ImportBatch).order_by(ImportBatch.id.desc()).limit(50)
        )
    ]
