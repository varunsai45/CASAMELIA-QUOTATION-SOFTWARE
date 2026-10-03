import base64, json, uuid, zipfile
from io import BytesIO
from pathlib import Path
from decimal import Decimal
import openpyxl, pymupdf, pytest
from sqlalchemy import select, func
from backend.app.current_import import import_current
from backend.app.models import (
    MasterProduct,
    SourcePrice,
    Area,
    AuditLog,
    QuotationVersion,
    ImportBatch,
    PriceConflict,
)
from .test_workflow import env, login, H, data


@pytest.fixture
def october(env):
    c, f = env
    with f() as db:
        import_current(db)
    return c, f


def kitchen_product(c):
    a = next(a for a in c.get("/areas").json() if a["name"] == "Kitchen")
    p = next(
        p
        for p in c.get(
            "/master-list",
            params={"area_id": a["id"], "item": "Base Unit", "limit": 500},
        ).json()["items"]
        if p["specification"] == "BWP Ply + HDHMR + Laminate"
    )
    return a, p


def test_current_import_and_area_permissions(october):
    c, f = october
    login(c)
    a, p = kitchen_product(c)
    assert Decimal(p["price"]) == 1990
    assert c.get("/master-list").json()["total"] == 662
    assert len(c.get("/price-conflicts").json()) == 32
    assert len(c.get("/areas").json()) == 9
    with f() as db:
        assert db.scalar(select(func.count()).select_from(SourcePrice)) == 1145
        assert (
            db.scalar(
                select(func.count())
                .select_from(MasterProduct)
                .where(MasterProduct.catalogue == "legacy")
            )
            == 238
        )
    bedroom = next(a for a in c.get("/areas").json() if a["name"] == "Bedrooms")
    assert not any(
        p["item"] == "Vanities"
        for p in c.get(
            "/master-list", params={"area_id": bedroom["id"], "limit": 500}
        ).json()["items"]
    )
    login(c, "sales")
    for path in ["/areas", "/imports/preview", "/settings/company"]:
        method = c.put if path == "/settings/company" else c.post
        assert (
            method(
                path,
                json=(
                    {"name": "Forbidden"}
                    if path == "/areas"
                    else (
                        {"filename": "x.xlsx", "content_base64": "AA=="}
                        if path == "/imports/preview"
                        else {"text": "Forbidden"}
                    )
                ),
                headers=H,
            ).status_code
            == 403
        )
    d = data(p["id"])
    d["sections"][0]["area_id"] = bedroom["id"]
    assert (
        "not assigned"
        in c.post("/quotations/preview", json=d, headers=H).json()["issues"][0]
    )


def test_override_immutable_pdf_excel_terms_company(october):
    c, f = october
    login(c, "sales")
    a, p = kitchen_product(c)
    customer = c.post(
        "/customers",
        json={
            "name": "Current project customer",
            "address": "Bengaluru",
            "phone": "9000000000",
            "notes": "User entered",
        },
        headers=H,
    ).json()
    d = data(p["id"])
    d["customer_id"] = customer["id"]
    d["sections"][0]["area_id"] = a["id"]
    item = d["sections"][0]["items"][0]
    item["quotation_rate"] = "1850"
    item["override_reason"] = "Project-specific agreed rate"
    preview = c.post("/quotations/preview", json=d, headers=H).json()
    i = preview["sections"][0]["items"][0]
    assert Decimal(i["master_rate"]) == 1990 and Decimal(i["quotation_rate"]) == 1850
    assert Decimal(i["amount"]) == 40700 and Decimal(preview["total"]) == 48026
    q = c.post("/quotations", json=d, headers=H).json()
    qid = q["id"]
    generated = c.post(f"/quotations/{qid}/generate", headers=H)
    assert generated.status_code == 200, generated.text
    rev = generated.json()["revision"]
    pdf = c.get(f"/quotations/{qid}/pdf").content
    excel = c.get(f"/quotations/{qid}/excel").content
    assert pdf.startswith(b"%PDF") and excel.startswith(b"PK")
    wb = openpyxl.load_workbook(BytesIO(excel), data_only=False)
    s = wb["Quotation"]
    line = next(r for r in range(15, s.max_row) if s.cell(r, 2).value == "Base Unit")
    assert s.cell(line, 8).value == 1990 and s.cell(line, 9).value == 1850
    assert s.cell(line, 6).value == f"=D{line}*E{line}"
    assert s.cell(line, 13).value == f"=ROUND(F{line}*G{line}*L{line}+K{line},2)"
    assert not s.protection.sheet and str(s.page_setup.paperSize) == s.PAPERSIZE_A4
    saved = c.get(f"/quotations/{qid}").json()
    oldterms = saved["payload"]["terms"]
    oldcompany = saved["payload"]["company_text"]
    assert saved["payload"]["sections"][0]["items"][0]["rate_changed_by"] is not None
    login(c)
    fields = {
        k: p[k]
        for k in [
            "item",
            "carcass",
            "shutter",
            "finish",
            "specification",
            "category",
            "unit",
            "measurement_mode",
            "price",
            "other",
            "hardware",
            "pricing_area",
            "active",
            "area_ids",
        ]
    }
    fields["price"] = "2100"
    assert c.put(f"/master-list/{p['id']}", json=fields, headers=H).status_code == 200
    assert (
        c.put(
            "/settings/terms", json={"terms": ["New approved terms"]}, headers=H
        ).status_code
        == 200
    )
    assert (
        c.put(
            "/settings/company",
            json={
                "text": "New company header",
                "bank_details": "Approved test bank details",
            },
            headers=H,
        ).status_code
        == 200
    )
    assert c.get(f"/quotations/{qid}/pdf").content == pdf
    assert c.get(f"/quotations/{qid}/excel").content == excel
    old = c.get(f"/quotations/{qid}").json()["payload"]
    assert old["terms"] == oldterms and old["company_text"] == oldcompany
    edit = json.loads(json.dumps(d))
    edit["revision"] = saved["revision"]
    edit["sections"][0]["items"][0]["quantity"] = "2"
    assert c.put(f"/quotations/{qid}", json=edit, headers=H).status_code == 200
    assert c.post(f"/quotations/{qid}/generate", headers=H).status_code == 200
    assert c.get(f"/quotations/{qid}/pdf?revision={rev}").content == pdf
    assert c.get(f"/quotations/{qid}/excel?revision={rev}").content == excel
    assert len(c.get(f"/quotations/{qid}/versions").json()) == 2
    with f() as db:
        assert (
            db.scalar(
                select(AuditLog).where(AuditLog.action == "quotation_rate_override")
            )
            is not None
        )
    out = Path(__file__).resolve().parents[2] / "output" / "october"
    out.mkdir(parents=True, exist_ok=True)
    (out / "verified-quotation.pdf").write_bytes(pdf)
    (out / "verified-quotation.xlsx").write_bytes(excel)
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    for n, page in enumerate(doc, 1):
        page.get_pixmap(matrix=pymupdf.Matrix(1.7, 1.7)).save(
            out / f"verified-page-{n}.png"
        )
    assert "Current project customer" not in "\n".join(
        p.get_text() for p in doc
    )  # quotation customer is explicitly entered in d
    assert "QA Customer" in "\n".join(p.get_text() for p in doc)


def test_unpriced_custom_hardware_and_conflicts(october):
    c, f = october
    login(c, "sales")
    a, p = kitchen_product(c)
    d = data(p["id"])
    d["sections"][0]["area_id"] = a["id"]
    i = d["sections"][0]["items"][0]
    i["quotation_rate"] = "0"
    i["override_reason"] = "Included in project package"
    assert c.post("/quotations/preview", json=d, headers=H).json()["total"] == "0.00"
    i["override_reason"] = ""
    assert (
        "reason" in c.post("/quotations/preview", json=d, headers=H).json()["issues"][0]
    )
    i.update(
        product_id=None,
        custom=True,
        item_label="Custom storage",
        description="Customer specified work",
        quotation_rate="500",
        override_reason="Custom project item",
    )
    assert (
        c.post("/quotations/preview", json=d, headers=H).json()["sections"][0]["items"][
            0
        ]["amount"]
        == "11000.00"
    )
    i["hardware_amount"] = "68000"
    i["hardware_area"] = "30"
    i["carcass_rate"] = "1270"
    assert abs(
        Decimal(
            c.post("/quotations/preview", json=d, headers=H).json()["sections"][0][
                "items"
            ][0]["quotation_rate"]
        )
        - Decimal("3536.666666666666666666666667")
    ) < Decimal(".000000001")
    login(c)
    conflict = c.get("/price-conflicts").json()[0]
    i.update(
        product_id=conflict["product_id"],
        custom=False,
        hardware_amount=None,
        hardware_area=None,
        carcass_rate=None,
    )
    d["sections"][0]["area_id"] = None
    assert (
        "conflict"
        in c.post("/quotations/preview", json=d, headers=H).json()["issues"][0]
    )


def test_reviewed_import_duplicate_and_change_history(october):
    c, f = october
    login(c)
    path = Path(__file__).resolve().parents[1] / "data" / "Qoute_1-10-2026.xlsx"
    raw = path.read_bytes()
    r = c.post(
        "/imports/preview",
        json={"filename": path.name, "content_base64": base64.b64encode(raw).decode()},
        headers=H,
    )
    assert r.status_code == 201, r.text
    result = r.json()
    assert (
        len(result["report"]["conflicts"]) == 32 and not result["report"]["duplicates"]
    )
    assert c.post(f"/imports/{result['id']}/apply", headers=H).status_code == 200
    assert c.post(f"/imports/{result['id']}/apply", headers=H).status_code == 409
    assert len(c.get("/price-conflicts").json()) == 32
    a, p = kitchen_product(c)
    assert c.get(f"/master-list/{p['id']}/history").json()["sources"]
    preview = c.post(
        "/imports/preview",
        json={"filename": path.name, "content_base64": base64.b64encode(raw).decode()},
        headers=H,
    ).json()
    with f() as db:
        batch = db.get(ImportBatch, preview["id"])
        updated = json.loads(json.dumps(batch.data))
        row = next(
            row
            for row in updated["products"]
            if row["lookup_price"] is not None
            and row["price"] is not None
            and Decimal(row["lookup_price"]) != Decimal(row["price"])
        )
        original = db.scalar(
            select(PriceConflict).where(
                PriceConflict.product_id == row["existing_id"],
                PriceConflict.status == "unresolved",
            )
        )
        original_values = (original.visible_price, original.lookup_price)
        row["price"] = str(Decimal(row["price"]) + 1)
        batch.data = updated
        db.commit()
    rejected = c.post(f"/imports/{preview['id']}/apply", headers=H)
    assert rejected.status_code == 409 and "existing price conflict" in rejected.text
    with f() as db:
        conflict = db.scalar(
            select(PriceConflict).where(
                PriceConflict.product_id == row["existing_id"],
                PriceConflict.status == "unresolved",
            )
        )
        assert (conflict.visible_price, conflict.lookup_price) == original_values


def test_admin_area_assignments_and_finalization(october):
    c, _ = october
    login(c)
    _, product = kitchen_product(c)
    area = c.post(
        "/areas", json={"name": "Configured project area", "position": 100}, headers=H
    ).json()
    assigned = c.put(
        f"/areas/{area['id']}/products",
        json={"product_ids": [product["id"]]},
        headers=H,
    )
    assert assigned.status_code == 200
    assert [
        p["id"]
        for p in c.get("/master-list", params={"area_id": area["id"]}).json()["items"]
    ] == [product["id"]]
    d = data(product["id"])
    d["sections"][0]["area_id"] = area["id"]
    quote = c.post("/quotations", json=d, headers=H).json()
    assert c.post(f"/quotations/{quote['id']}/finalize", headers=H).status_code == 409
    assert c.post(f"/quotations/{quote['id']}/generate", headers=H).status_code == 200
    pdf = c.get(f"/quotations/{quote['id']}/pdf").content
    assert (
        c.post(f"/quotations/{quote['id']}/finalize", headers=H).json()["status"]
        == "finalized"
    )
    assert c.get(f"/quotations/{quote['id']}/pdf").content == pdf
    login(c, "sales")
    assert (
        c.put(
            f"/areas/{area['id']}/products", json={"product_ids": []}, headers=H
        ).status_code
        == 403
    )
    assert c.post(f"/quotations/{quote['id']}/finalize", headers=H).status_code == 403


def test_excel_formula_injection_is_text(october):
    from backend.app.excel import generate_excel

    c, f = october
    login(c)
    a, p = kitchen_product(c)
    d = data(p["id"])
    d["customer_name"] = "=1+2"
    d["sections"][0]["items"][0]["description"] = '=HYPERLINK("https://example.com")'
    q = c.post("/quotations", json=d, headers=H).json()
    c.post(f"/quotations/{q['id']}/generate", headers=H)
    wb = openpyxl.load_workbook(BytesIO(c.get(f"/quotations/{q['id']}/excel").content))
    assert wb["Quotation"]["C9"].data_type == "s"
    assert wb["Quotation"]["C16"].data_type == "s"


def test_equal_rate_override_and_reset(october):
    c, factory = october
    login(c, "sales")
    a, product = kitchen_product(c)
    with factory() as db:
        db.get(MasterProduct, product["id"]).price = Decimal("100")
        db.commit()
    d = data(product["id"])
    d["sections"][0]["area_id"] = a["id"]
    line = d["sections"][0]["items"][0]
    line["quotation_rate"] = "100.00"
    line["override_reason"] = "Stale reason"
    preview = c.post("/quotations/preview", json=d, headers=H).json()
    result = preview["sections"][0]["items"][0]
    assert result["rate_override"] is False and result["override_reason"] is None
    assert result["rate_changed_by"] is None and result["rate_changed_at"] is None
    line["quotation_rate"] = "90"
    line["override_reason"] = None
    result = c.post("/quotations/preview", json=d, headers=H).json()["sections"][0][
        "items"
    ][0]
    assert result["rate_override"] is True and "reason" in result["issue"]
    line["override_reason"] = "Customer-specific pricing"
    quote = c.post("/quotations", json=d, headers=H).json()
    result = quote["payload"]["sections"][0]["items"][0]
    assert result["rate_override"] is True and result["rate_changed_by"] is not None
    line["quotation_rate"] = "100"
    d["revision"] = quote["revision"]
    updated = c.put(f"/quotations/{quote['id']}", json=d, headers=H).json()
    result = updated["payload"]["sections"][0]["items"][0]
    assert result["rate_override"] is False and result["override_reason"] is None
    assert c.post(f"/quotations/{quote['id']}/generate", headers=H).status_code == 200


def test_customer_history_uses_saved_documents_and_ownership(october):
    c, _ = october
    login(c, "sales")
    a, p = kitchen_product(c)
    customer = c.post(
        "/customers",
        json={"name": "Current V2 customer", "address": "Current address"},
        headers=H,
    ).json()
    d = data(p["id"])
    d["customer_id"] = customer["id"]
    d["customer_name"] = customer["name"]
    d["sections"][0]["area_id"] = a["id"]
    quote = c.post("/quotations", json=d, headers=H).json()
    generated = c.post(f"/quotations/{quote['id']}/generate", headers=H)
    assert generated.status_code == 200 and generated.headers[
        "content-type"
    ].startswith("application/json")
    pdf = c.get(f"/quotations/{quote['id']}/pdf?inline=true").content
    excel = c.get(f"/quotations/{quote['id']}/excel").content
    history = c.get(f"/customers/{customer['id']}").json()
    assert history["quotations"][0]["id"] == quote["id"] and history["projects"] == [
        d["project_reference"]
    ]
    assert c.get("/customers").json()[0]["quotation_count"] == 1
    assert c.get("/quotations", params={"status": "draft"}).json()["total"] == 0
    assert c.get("/quotations", params={"status": "generated"}).json()["total"] == 1
    assert (
        c.get(f"/quotations/{quote['id']}/versions").json()[0]["generated_by"]
        == "sales"
    )
    assert c.get(f"/quotations/{quote['id']}/pdf").content == pdf
    assert c.get(f"/quotations/{quote['id']}/excel").content == excel
    login(c)
    assert c.get(f"/customers/{customer['id']}").status_code == 200
    private = c.post(
        "/customers", json={"name": "Admin-only customer"}, headers=H
    ).json()
    login(c, "sales")
    assert c.get(f"/customers/{private['id']}").status_code == 404
